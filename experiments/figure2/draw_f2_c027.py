from pathlib import Path
import html
import math
import xml.etree.ElementTree as ET

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Rectangle, PathPatch
from matplotlib.path import Path as MplPath

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
W, H = 1412, 587

# Stage A uses neutral grayscale; Stage B freezes every coordinate below and
# applies only the palette/typography/line treatment from experiments/README.md.
STYLES = {
    "A": {
        "font": "DejaVu Sans", "text": "#202020", "edge": "#202020",
        "dash": "#202020", "token_group": "#F1F1F1", "channel_group": "#F7F7F7",
        "box": "#E9E9E9", "memory": "#FAFAFA", "slot_a": "#E2E2E2",
        "slot_b": "#D7D7D7", "slot_c": "#EEEEEE", "slot_d": "#C9C9C9",
    },
    "B": {
        "font": "DejaVu Sans", "text": "#303330", "edge": "#78838F",
        "dash": "#78838F", "token_group": "#EAF3F8", "channel_group": "#F0EFF8",
        "box": "#E3F2EB", "memory": "#F7F7F4", "slot_a": "#D9EEDB",
        "slot_b": "#F2D2CE", "slot_c": "#F8E7BF", "slot_d": "#E8D7E4",
        "green": "#3AA278", "coral": "#D75C5D", "periwinkle": "#6667AB",
        "blue": "#4395C4", "amber": "#D99532", "berry": "#B65C91",
        "slate": "#78838F",
    },
}

# All geometry is in the historical 1412x587 raster coordinate system (origin
# at the top-left). Keeping these coordinates fixed is the A->B gate.
GROUPS = [
    ("decode_kernel", 163, 41, 906, 504, "dashed"),
    ("token_mixer", 303, 81, 746, 202, "rounded"),
    ("channel_mixer", 303, 322, 746, 204, "rounded"),
]

BOXES = [
    ("ps_host", 61, 243, 81, 119, "PS\nHost", "box"),
    ("embed", 202, 143, 81, 119, "Embed", "box"),
    ("lm_head", 202, 383, 81, 121, "LM\nHead", "box"),
    ("qkvg", 323, 142, 162, 121, "INT8\nMAC Engine\nQKVG", "blue"),
    ("conv", 524, 142, 183, 121, "FP32 Engine\nConv / L2\nGatedDeltaNet", "green"),
    ("o_proj", 747, 142, 159, 121, "INT8\nMAC Engine\nO", "blue"),
    ("add_top", 947, 142, 61, 121, "Add", "amber"),
    ("rms", 947, 383, 61, 121, "RMS\nNorm", "amber"),
    ("w13", 766, 383, 161, 121, "INT8\nMAC Engine\nW1 / W3", "blue"),
    ("silu", 585, 383, 162, 121, "FP32 Engine\nSiLu ×", "green"),
    ("w2", 404, 383, 162, 121, "INT8\nMAC Engine\nW2", "blue"),
    ("add_bottom", 323, 383, 61, 121, "Add", "amber"),
    ("uram", 1128, 41, 223, 222, "URAM", "memory"),
    ("bram", 1128, 303, 223, 242, "BRAM", "memory"),
    ("uram_s1", 1149, 101, 80, 62, "S₁", "slot_a"),
    ("uram_s2", 1249, 101, 80, 62, "S₂", "slot_c"),
    ("uram_ellipsis", 1149, 181, 80, 62, "...", "slot_c"),
    ("uram_st", 1249, 181, 80, 62, "Sₜ", "slot_b"),
    ("bram_conv", 1149, 363, 181, 39, "Conv History", "slot_a"),
    ("bram_params", 1149, 423, 181, 40, "Layer Params", "slot_c"),
    ("bram_buffers", 1149, 484, 181, 40, "Buffers", "slot_d"),
]

LOOP_BEZIER_SEGMENTS = [
    ((1068, 136), (1082, 123), (1107, 112), (1107, 80)),
    ((1107, 80), (1107, 49), (1083, 28), (1055, 28)),
    ((1055, 28), (1024, 28), (1007, 45), (1009, 68)),
    ((1009, 68), (1011, 95), (1037, 117), (1058, 121)),
]


def sample_loop_points(per_segment=30):
    points = []
    for seg_idx, (p0, p1, p2, p3) in enumerate(LOOP_BEZIER_SEGMENTS):
        for k in range(per_segment + 1):
            if seg_idx and k == 0:
                continue
            t = k / per_segment
            points.append(tuple(round((1 - t) ** 3 * p0[i] + 3 * (1 - t) ** 2 * t * p1[i] + 3 * (1 - t) * t * t * p2[i] + t ** 3 * p3[i], 2) for i in (0, 1)))
    return points

ARROWS = [
    ("input_tokens", [(101, 42), (101, 242)], "Input Tokens", "vertical"),
    ("output_tokens", [(101, 362), (101, 544)], "Output Tokens", "vertical"),
    ("host_to_embed", [(141, 272), (242, 272), (242, 263)], "", "orthogonal"),
    ("lm_to_host", [(242, 383), (242, 333), (142, 333)], "", "orthogonal"),
    ("embed_to_qkvg", [(283, 202), (323, 202)], "", "horizontal"),
    ("qkvg_to_conv", [(485, 202), (524, 202)], "", "horizontal"),
    ("conv_to_o", [(707, 202), (747, 202)], "", "horizontal"),
    ("o_to_add", [(906, 202), (947, 202)], "", "horizontal"),
    ("add_to_rms", [(978, 263), (978, 383)], "", "vertical"),
    ("rms_to_w13", [(947, 444), (927, 444)], "", "horizontal"),
    ("w13_to_silu", [(766, 444), (747, 444)], "", "horizontal"),
    ("silu_to_w2", [(585, 444), (566, 444)], "", "horizontal"),
    ("w2_to_add", [(404, 444), (384, 444)], "", "horizontal"),
    ("add_to_lm", [(323, 444), (283, 444)], "", "horizontal"),
    ("uram_link", [(1068, 162), (1128, 162)], "", "horizontal"),
    ("bram_link", [(1068, 423), (1128, 423)], "", "horizontal"),
    ("loop_back", sample_loop_points(), "×T", "loop"),
]


def palette(stage):
    return STYLES[stage]


def fill_color(stage, role):
    s = palette(stage)
    if stage == "A":
        return s[role] if role in s else s["box"]
    if role == "blue":
        return "#DDECF4"
    if role == "green":
        return "#D9EEDB"
    if role == "amber":
        return "#F8E7BF"
    if role == "box":
        return "#F3F3F0"
    return s.get(role, s["box"])


def text_size(role, stage):
    if role in {"decode_kernel", "token_mixer", "channel_mixer"}:
        return 13 if stage == "B" else 12
    return 10 if stage == "B" else 10


def draw_box(ax, stage, ident, x, y, w, h, text, role):
    s = palette(stage)
    if ident in {"uram", "bram"}:
        rounding = 0.04
    else:
        rounding = 0.025
    patch = FancyBboxPatch((x, y), w, h, boxstyle=f"round,pad=0,rounding_size={rounding * min(w, h)}",
                           facecolor=fill_color(stage, role), edgecolor=s["edge"],
                           linewidth=1.1 if stage == "A" else 0.9)
    ax.add_patch(patch)
    if ident in {"uram", "bram"}:
        ax.text(x + w / 2, y + 27, text, ha="center", va="center", family=s["font"],
                fontsize=12 if stage == "A" else 11, color=s["text"])
        return
    lines = text.split("\n")
    fontsize = 10 if ident not in {"uram", "bram"} else 10
    if ident in {"uram_s1", "uram_s2", "uram_ellipsis", "uram_st"}:
        fontsize = 11
    ax.text(x + w / 2, y + h / 2, "\n".join(lines), ha="center", va="center",
            family=s["font"], fontsize=fontsize, color=s["text"], linespacing=1.08)


def draw_group(ax, stage, ident, x, y, w, h, kind):
    s = palette(stage)
    if kind == "dashed":
        ax.add_patch(Rectangle((x, y), w, h, fill=False, edgecolor=s["dash"],
                               linewidth=1.0, linestyle=(0, (8, 7))))
        ax.text(x + w / 2, y + 24, "Decode Kernel", ha="center", va="center",
                family=s["font"], fontsize=13, color=s["text"])
    else:
        role = "token_group" if ident == "token_mixer" else "channel_group"
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0,rounding_size=22",
                                    facecolor=s[role], edgecolor=s["edge"], linewidth=1.0))
        title = "Token Mixer" if ident == "token_mixer" else "Channel Mixer"
        ax.text(x + w / 2, y + 28, title, ha="center", va="center", family=s["font"],
                fontsize=13, color=s["text"])


def arrow_segment(ax, p1, p2, color, width, head=True):
    style = "-|>" if head else "-"
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle=style, mutation_scale=12,
                                 linewidth=width, color=color, shrinkA=0, shrinkB=0))


def draw_arrows_mpl(ax, stage):
    s = palette(stage)
    width = 1.15 if stage == "A" else 1.05
    for ident, points, label, kind in ARROWS:
        if ident == "loop_back":
            verts = [(1068, 136), (1082, 123), (1107, 112), (1107, 80),
                     (1107, 49), (1083, 28), (1055, 28), (1024, 28),
                     (1007, 45), (1009, 68), (1011, 95), (1037, 117), (1058, 121)]
            codes = [MplPath.MOVETO] + [MplPath.CURVE4] * (len(verts) - 1)
            ax.add_patch(PathPatch(MplPath(verts, codes), fill=False, color=s["edge"], linewidth=width))
            arrow_segment(ax, (1037, 117), (1058, 121), s["edge"], width, head=True)
            ax.text(1078, 49, label, ha="center", va="center", family=s["font"],
                    fontsize=10, color=s["text"])
            continue
        for idx in range(len(points) - 1):
            arrow_segment(ax, points[idx], points[idx + 1], s["edge"], width,
                          head=(idx == len(points) - 2))
        if ident == "input_tokens":
            ax.text(78, 142, label, rotation=90, ha="center", va="center",
                    family=s["font"], fontsize=10, color=s["text"])
        elif ident == "output_tokens":
            ax.text(78, 451, label, rotation=90, ha="center", va="center",
                    family=s["font"], fontsize=10, color=s["text"])
    ax.text(124, 147, "Who is", rotation=90, ha="center", va="center",
            family=s["font"], fontsize=10, color=s["text"])
    ax.text(124, 442, "Adam", rotation=90, ha="center", va="center",
            family=s["font"], fontsize=10, color=s["text"])


def render_matplotlib(stage, out_dir):
    s = palette(stage)
    out_dir.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(12, 12 * H / W))
    ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")
    fig.patch.set_facecolor("white"); ax.set_facecolor("white")
    for ident, x, y, w, h, kind in GROUPS:
        draw_group(ax, stage, ident, x, y, w, h, kind)
    for ident, x, y, w, h, text, role in BOXES:
        draw_box(ax, stage, ident, x, y, w, h, text, role)
    draw_arrows_mpl(ax, stage)
    fig.subplots_adjust(left=0, right=1, bottom=0, top=1)
    fig.savefig(out_dir / f"figure2_stage_{stage}.pdf", facecolor="white", dpi=300)
    fig.savefig(out_dir / f"figure2_stage_{stage}.png", facecolor="white", dpi=300)
    plt.close(fig)


def svg_text(parent, stage, x, y, text, size=14, anchor="middle", weight="normal", rotate=None):
    s = palette(stage)
    attrs = {"x": str(x), "y": str(y), "font-size": str(size), "font-family": s["font"],
             "fill": s["text"], "text-anchor": anchor, "font-weight": weight}
    if rotate is not None:
        attrs["transform"] = f"rotate({rotate} {x} {y})"
    node = ET.SubElement(parent, "text", attrs)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        t = ET.SubElement(node, "tspan", {"x": str(x), "dy": "0" if i == 0 else str(size * 1.08)})
        t.text = line
    return node


def svg_rect(parent, stage, ident, x, y, w, h, fill, radius=0, dashed=False, stroke_width=1.1):
    s = palette(stage)
    attrs = {"id": ident, "x": str(x), "y": str(y), "width": str(w), "height": str(h),
             "fill": fill, "stroke": s["edge"], "stroke-width": str(stroke_width)}
    if radius:
        attrs["rx"] = str(radius); attrs["ry"] = str(radius)
    if dashed:
        attrs["fill"] = "none"; attrs["stroke"] = s["dash"]; attrs["stroke-dasharray"] = "10,9"
    return ET.SubElement(parent, "rect", attrs)


def draw_groups_svg(root, stage):
    s = palette(stage)
    for ident, x, y, w, h, kind in GROUPS:
        g = ET.SubElement(root, "g", {"id": f"group_{ident}"})
        if kind == "dashed":
            svg_rect(g, stage, f"boundary_{ident}", x, y, w, h, "none", dashed=True)
            svg_text(g, stage, x + w / 2, y + 28, "Decode Kernel", 18 if stage == "A" else 17,
                     weight="600")
        else:
            fill = s["token_group" if ident == "token_mixer" else "channel_group"]
            svg_rect(g, stage, f"panel_{ident}", x, y, w, h, fill, radius=22, stroke_width=1.0)
            title = "Token Mixer" if ident == "token_mixer" else "Channel Mixer"
            svg_text(g, stage, x + w / 2, y + 29, title, 18 if stage == "A" else 17)


def draw_boxes_svg(root, stage):
    s = palette(stage)
    for ident, x, y, w, h, text, role in BOXES:
        g = ET.SubElement(root, "g", {"id": f"box_{ident}"})
        radius = 20 if ident in {"uram", "bram"} else 8
        svg_rect(g, stage, f"rect_{ident}", x, y, w, h, fill_color(stage, role), radius=radius,
                 stroke_width=1.1 if stage == "A" else 0.9)
        if ident in {"uram", "bram"}:
            svg_text(g, stage, x + w / 2, y + 28, text, 17 if stage == "A" else 15, weight="600")
        else:
            fs = 15 if stage == "A" else 14
            if ident in {"uram_s1", "uram_s2", "uram_ellipsis", "uram_st"}:
                fs = 17 if stage == "A" else 15
            lines = text.split("\n")
            top = y + h / 2 - (len(lines) - 1) * fs * .52
            for idx, line in enumerate(lines):
                svg_text(g, stage, x + w / 2, top + idx * fs * 1.08, line, fs)


def draw_arrows_svg(root, stage):
    s = palette(stage)
    defs = root.find("defs")
    marker = ET.SubElement(defs, "marker", {"id": f"arrow_{stage}", "markerWidth": "9",
                                              "markerHeight": "9", "refX": "8", "refY": "4.5",
                                              "orient": "auto"})
    ET.SubElement(marker, "path", {"d": "M0,0 L9,4.5 L0,9 Z", "fill": s["edge"]})
    for ident, points, label, kind in ARROWS:
        g = ET.SubElement(root, "g", {"id": f"flow_{ident}"})
        if ident == "loop_back":
            d = "M1068,136 C1082,123 1107,112 1107,80 C1107,49 1083,28 1055,28 C1024,28 1007,45 1009,68 C1011,95 1037,117 1058,121"
            ET.SubElement(g, "path", {"d": d, "fill": "none", "stroke": s["edge"],
                                        "stroke-width": "1.2", "marker-end": f"url(#arrow_{stage})"})
            svg_text(g, stage, 1078, 49, label, 15 if stage == "A" else 14)
            continue
        for idx in range(len(points) - 1):
            x1, y1 = points[idx]; x2, y2 = points[idx + 1]
            attrs = {"x1": str(x1), "y1": str(y1), "x2": str(x2), "y2": str(y2),
                     "stroke": s["edge"], "stroke-width": "1.2", "fill": "none"}
            if idx == len(points) - 2:
                attrs["marker-end"] = f"url(#arrow_{stage})"
            ET.SubElement(g, "line", attrs)
        if ident == "input_tokens":
            svg_text(g, stage, 78, 142, label, 15 if stage == "A" else 14, rotate=-90)
        elif ident == "output_tokens":
            svg_text(g, stage, 78, 451, label, 15 if stage == "A" else 14, rotate=-90)
    svg_text(root, stage, 124, 147, "Who is", 15 if stage == "A" else 14, rotate=-90)
    svg_text(root, stage, 124, 442, "Adam", 15 if stage == "A" else 14, rotate=-90)


def render_svg(stage, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    root = ET.Element("svg", {"xmlns": "http://www.w3.org/2000/svg", "width": str(W),
                               "height": str(H), "viewBox": f"0 0 {W} {H}"})
    ET.SubElement(root, "rect", {"x": "0", "y": "0", "width": str(W), "height": str(H),
                                  "fill": "#FFFFFF"})
    ET.SubElement(root, "defs")
    draw_groups_svg(root, stage)
    draw_boxes_svg(root, stage)
    draw_arrows_svg(root, stage)
    ET.ElementTree(root).write(out_dir / f"figure2_stage_{stage}.svg", encoding="utf-8", xml_declaration=True)


def ppt_rgb(value):
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def ppt_text(shape, stage, text, size=10, bold=False, rotate=0, wrap=True):
    s = palette(stage)
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = wrap; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    for idx, line in enumerate(text.split("\n")):
        if idx:
            p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = line
        run.font.name = s["font"]; run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = ppt_rgb(s["text"])
    shape.rotation = rotate


def clear_ppt_style(shape):
    """Remove the theme style/effect reference so Impress cannot add a shadow."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def ppt_arrowhead(shape):
    from pptx.oxml.xmlchemy import OxmlElement
    ln = shape.line._get_or_add_ln()
    tail = OxmlElement("a:tailEnd")
    tail.set("type", "triangle"); tail.set("w", "sm"); tail.set("len", "sm")
    ln.append(tail)


def ppt_freeform_loop(slide, stage, sx, sy):
    s = palette(stage)
    x_scale = Inches(13.333) / W
    y_scale = Inches(13.333 * H / W) / H
    points = sample_loop_points(30)
    builder = slide.shapes.build_freeform(points[0][0], points[0][1], scale=(x_scale, y_scale))
    builder.add_line_segments(points[1:], close=False)
    shape = builder.convert_to_shape()
    clear_ppt_style(shape)
    shape.name = "flow_loop_back_freeform"
    shape.fill.background()
    shape.line.color.rgb = ppt_rgb(s["edge"])
    shape.line.width = Pt(0.8 if stage == "A" else 0.7)
    ppt_arrowhead(shape)
    return shape


def ppt_box(slide, stage, ident, x, y, w, h, text, role, sx, sy):
    s = palette(stage)
    if ident in {"uram", "bram"}:
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x * sx), Inches(y * sy),
                                     Inches(w * sx), Inches(h * sy))
        clear_ppt_style(sh)
        sh.name = f"box_{ident}"; sh.fill.solid(); sh.fill.fore_color.rgb = ppt_rgb(fill_color(stage, role))
        sh.line.color.rgb = ppt_rgb(s["edge"]); sh.line.width = Pt(0.9)
        lab = slide.shapes.add_textbox(Inches((x + 20) * sx), Inches((y + 9) * sy), Inches((w - 40) * sx), Inches(.30))
        clear_ppt_style(lab)
        ppt_text(lab, stage, text, 12 if stage == "A" else 11, bold=(stage == "B"))
        return
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x * sx), Inches(y * sy),
                                Inches(w * sx), Inches(h * sy))
    clear_ppt_style(sh)
    sh.name = f"box_{ident}"; sh.fill.solid(); sh.fill.fore_color.rgb = ppt_rgb(fill_color(stage, role))
    sh.line.color.rgb = ppt_rgb(s["edge"]); sh.line.width = Pt(0.9)
    size = 10 if stage == "A" else 9.5
    if ident in {"uram_s1", "uram_s2", "uram_ellipsis", "uram_st"}:
        size = 11 if stage == "A" else 10
    if ident in {"add_top", "rms", "add_bottom"}:
        size = 9 if stage == "A" else 8.5
    ppt_text(sh, stage, text, size)


def ppt_group(slide, stage, ident, x, y, w, h, kind, sx, sy):
    s = palette(stage)
    if kind == "dashed":
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x * sx), Inches(y * sy), Inches(w * sx), Inches(h * sy))
        clear_ppt_style(sh)
        sh.name = f"group_{ident}"; sh.fill.background(); sh.line.color.rgb = ppt_rgb(s["dash"])
        sh.line.width = Pt(0.8); sh.line.dash_style = MSO_LINE.DASH
        lab = slide.shapes.add_textbox(Inches((x + w / 2 - 90) * sx), Inches((y + 11) * sy), Inches(180 * sx), Inches(.25))
        clear_ppt_style(lab); lab.name = "label_decode_kernel"; ppt_text(lab, stage, "Decode Kernel", 12 if stage == "A" else 11, bold=(stage == "B"))
    else:
        role = "token_group" if ident == "token_mixer" else "channel_group"
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x * sx), Inches(y * sy), Inches(w * sx), Inches(h * sy))
        clear_ppt_style(sh)
        sh.name = f"group_{ident}"; sh.fill.solid(); sh.fill.fore_color.rgb = ppt_rgb(s[role])
        sh.line.color.rgb = ppt_rgb(s["edge"]); sh.line.width = Pt(0.8)
        title = "Token Mixer" if ident == "token_mixer" else "Channel Mixer"
        lab = slide.shapes.add_textbox(Inches((x + w / 2 - 110) * sx), Inches((y + 14) * sy), Inches(220 * sx), Inches(.26))
        clear_ppt_style(lab); lab.name = f"label_{ident}"; ppt_text(lab, stage, title, 12 if stage == "A" else 11)


def render_pptx(stage, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(13.333 * H / W)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sx = 13.333 / W; sy = (13.333 * H / W) / H
    for ident, x, y, w, h, kind in GROUPS:
        ppt_group(slide, stage, ident, x, y, w, h, kind, sx, sy)
    for ident, x, y, w, h, text, role in BOXES:
        ppt_box(slide, stage, ident, x, y, w, h, text, role, sx, sy)
    s = palette(stage); width = 0.8 if stage == "A" else 0.7
    for ident, points, label, kind in ARROWS:
        if ident == "loop_back":
            ppt_freeform_loop(slide, stage, sx, sy)
            lab = slide.shapes.add_textbox(Inches(1060 * sx), Inches(38 * sy), Inches(36 * sx), Inches(22 * sy))
            clear_ppt_style(lab); lab.name = "label_loop_back"; ppt_text(lab, stage, label, 10 if stage == "A" else 9, wrap=False)
            continue
        for idx in range(len(points) - 1):
            x1, y1 = points[idx]; x2, y2 = points[idx + 1]
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1 * sx), Inches(y1 * sy),
                                               Inches(x2 * sx), Inches(y2 * sy))
            clear_ppt_style(line)
            line.name = f"flow_{ident}_{idx}"; line.line.color.rgb = ppt_rgb(s["edge"]); line.line.width = Pt(width)
            if idx == len(points) - 2:
                ppt_arrowhead(line)
        if ident == "input_tokens":
            lab = slide.shapes.add_textbox(Inches(28 * sx), Inches(131 * sy), Inches(100 * sx), Inches(22 * sy)); clear_ppt_style(lab); lab.name = "label_input_tokens"; ppt_text(lab, stage, label, 10 if stage == "A" else 9, rotate=270, wrap=False)
        elif ident == "output_tokens":
            lab = slide.shapes.add_textbox(Inches(28 * sx), Inches(440 * sy), Inches(100 * sx), Inches(22 * sy)); clear_ppt_style(lab); lab.name = "label_output_tokens"; ppt_text(lab, stage, label, 10 if stage == "A" else 9, rotate=270, wrap=False)
    for text, y in [("Who is", 147), ("Adam", 442)]:
        width = 60 if text == "Who is" else 45
        lab = slide.shapes.add_textbox(Inches((124 - width / 2) * sx), Inches((y - 11) * sy), Inches(width * sx), Inches(22 * sy)); clear_ppt_style(lab); lab.name = f"label_{text.replace(' ', '_')}"; ppt_text(lab, stage, text, 10 if stage == "A" else 9, rotate=270, wrap=False)
    prs.save(out_dir / f"figure2_stage_{stage}.pptx")
